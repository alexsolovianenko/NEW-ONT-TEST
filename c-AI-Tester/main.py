import boto3
import openai
import os
from dotenv import load_dotenv
import logging
from datetime import datetime
from docx import Document
from fpdf import FPDF, HTMLMixin
from adobe.pdfservices.operation.auth.service_principal_credentials import ServicePrincipalCredentials
from adobe.pdfservices.operation.exception.exceptions import ServiceApiException, ServiceUsageException, SdkException
from adobe.pdfservices.operation.io.cloud_asset import CloudAsset
from adobe.pdfservices.operation.io.stream_asset import StreamAsset
from adobe.pdfservices.operation.pdf_services import PDFServices
from adobe.pdfservices.operation.pdf_services_media_type import PDFServicesMediaType
from adobe.pdfservices.operation.pdfjobs.jobs.export_pdf_job import ExportPDFJob
from adobe.pdfservices.operation.pdfjobs.params.export_pdf.export_pdf_params import ExportPDFParams
from adobe.pdfservices.operation.pdfjobs.params.export_pdf.export_pdf_target_format import ExportPDFTargetFormat
from adobe.pdfservices.operation.pdfjobs.params.export_pdf.export_ocr_locale import ExportOCRLocale
from adobe.pdfservices.operation.pdfjobs.result.export_pdf_result import ExportPDFResult

load_dotenv()
first_aws_access_key_id = os.getenv('FIRST_AWS_ACCESS_KEY_ID')  
first_aws_secret_access_key = os.getenv('FIRST_AWS_SECRET_ACCESS_KEY')
first_aws_region = os.getenv('FIRST_AWS_REGION')  
first_s3_bucket_name = os.getenv('FIRST_S3_BUCKET_NAME')  
second_aws_access_key_id = os.getenv('SECOND_AWS_ACCESS_KEY_ID')
second_aws_secret_access_key = os.getenv('SECOND_AWS_SECRET_ACCESS_KEY')
second_aws_region = os.getenv('SECOND_AWS_REGION')
second_s3_bucket_name = os.getenv('SECOND_S3_BUCKET_NAME')
openai.api_key = os.getenv('OPENAI_API_KEY')

s3_client = boto3.client(
    's3',
    aws_access_key_id=first_aws_access_key_id,
    aws_secret_access_key=first_aws_secret_access_key
)
response = s3_client.list_objects_v2(Bucket=first_s3_bucket_name)

if 'Contents' in response:
    sorted_objects = sorted(response['Contents'], key=lambda obj: obj['LastModified'], reverse=True)
    latest_file_metadata = sorted_objects[0]
    object_key = latest_file_metadata['Key']

    print(f"Fetching the latest file: {object_key}")

    file_response = s3_client.get_object(Bucket=first_s3_bucket_name, Key=object_key)
    file_content = file_response['Body'].read()

    with open(object_key, "wb") as file:
        file.write(file_content)

    print(f"File downloaded successfully as '{object_key}'.")
else:
    print("No files found in the S3 bucket.")

logging.basicConfig(level=logging.INFO)

class ExportPDFToDOCXWithOCROption:
    def __init__(self, file_path):
        try:
            global docx_output_file_path
            with open(file_path, 'rb') as file:
                input_stream = file.read()

            # Retrieve environment variables for PDF Services credentials
            pdf_services_client_id = os.getenv('PDF_SERVICES_CLIENT_ID')
            pdf_services_client_secret = os.getenv('PDF_SERVICES_CLIENT_SECRET')

            if not pdf_services_client_id or not pdf_services_client_secret:
                raise EnvironmentError("Missing PDF Services credentials: 'PDF_SERVICES_CLIENT_ID' or 'PDF_SERVICES_CLIENT_SECRET'")

            credentials = ServicePrincipalCredentials(
                client_id=pdf_services_client_id,
                client_secret=pdf_services_client_secret
            )

            pdf_services = PDFServices(credentials=credentials)

            input_asset = pdf_services.upload(input_stream=input_stream, mime_type=PDFServicesMediaType.PDF)

            export_pdf_params = ExportPDFParams(target_format=ExportPDFTargetFormat.DOCX, ocr_lang=ExportOCRLocale.EN_US)

            export_pdf_job = ExportPDFJob(input_asset=input_asset, export_pdf_params=export_pdf_params)

            location = pdf_services.submit(export_pdf_job)
            pdf_services_response = pdf_services.get_job_result(location, ExportPDFResult)

            result_asset: CloudAsset = pdf_services_response.get_result().get_asset()
            stream_asset: StreamAsset = pdf_services.get_content(result_asset)

            timestamp = datetime.now().strftime("%Y%m%d%H%M%S")
            docx_output_file_path = f'output/ExportPDFToDOCXWithOCROption/export_{timestamp}.docx'
            os.makedirs(os.path.dirname(docx_output_file_path), exist_ok=True)

            with open(docx_output_file_path, "wb") as file:
                file.write(stream_asset.get_input_stream())

            print(f"PDF successfully converted to DOCX. File saved as '{docx_output_file_path}'.")

        except (ServiceApiException, ServiceUsageException, SdkException) as e:
            logging.exception(f'Exception encountered while executing operation: {e}')

def extract_questions_from_docx(docx_file_path):
    try:
        document = Document(docx_file_path)
        docx_content = "\n".join([paragraph.text for paragraph in document.paragraphs])

        prompt = f"""
You are an AI tutor tasked with creating a structured practice test based on the provided educational material. The test should follow the Ontario curriculum format and include the following sections:

1. **Knowledge**: 5 questions that test factual recall and understanding of the material. Each question is accompanied by 3-4 multiple choice using letters a), b)...
2. **Thinking**: 5 questions that require critical thinking and problem-solving skills.
3. **Application**: 5 questions that require applying knowledge to new situations or solving practical problems.
4. **Communication**: 5 questions that assess the ability to explain concepts clearly and effectively.

- Ensure the questions are relevant to the content provided in the document.
- Rewrite the questions to align with the curriculum but maintain the core concepts.
- If the document references images or diagrams, include placeholders for them (e.g., "[Refer to Figure 1]") and ensure the questions reference these appropriately.
- Provide clear and accurate answers for each question.
- Format the output as a structured test with numbered questions and labeled sections.

- After all the question, keep in mind which is the right answer and write it under the **Solutions** tab.
  - 

Here is the content of the document:
{docx_content}
"""
        # openai API runs w/ instructions
        response = openai.chat.completions.create(
            model="gpt-4",
            messages=[
                {"role": "system", "content": "You are a high school teacher creating a structured practice test with complete answers."},
                {"role": "user", "content": prompt}
            ],
            max_tokens=2500,
            temperature=0.5
        )
       
        return response.choices[0].message.content
    except openai.error.OpenAIError as e:
        logging.error(f"Error extracting questions with OpenAI: {e}")
        return "Error: Unable to process the request."
    except Exception as e:
        logging.error(f"Error reading DOCX file: {e}")
        return "Error: Unable to read the DOCX file."

class StyledPDF(FPDF, HTMLMixin):
    def header(self):
        self.set_fill_color(91, 44, 250)  
        self.set_text_color(255, 255, 255)  
        self.set_font("Arial", style="B", size=10)  # Use default font
        self.set_text_color(0, 0, 0)
        self.set_xy(-40, 5)
        self.cell(30, 10, f"Page {self.page_no()}", align="R")

        # Title
        self.set_xy(10, 18)
        self.set_text_color(255, 255, 255)
        self.set_font("Arial", style="B", size=46)  # Use default font
        self.cell(0, 28, "Ontario Tests", ln=True, align="C", fill=True)
        self.ln(5)

    def footer(self):
        from datetime import datetime
        self.set_y(-15)
        self.set_font("Arial", style="", size=11)  # Use default font
        self.set_text_color(138, 153, 163)  # #8a99a3
        dev_notice = f"© {datetime.now().year} Ontario Tests | Early Development Version"
        self.cell(0, 10, dev_notice, align="C")
        self.set_text_color(150, 150, 150)

if __name__ == "__main__":
    ExportPDFToDOCXWithOCROption(object_key)

    questions = extract_questions_from_docx(docx_output_file_path)
    print("Generated Practice Test:", questions)

    # name file
    file_name_parts = os.path.splitext(object_key)[0].split("_")
    subject = file_name_parts[0] if len(file_name_parts) > 0 else "Subject"
    grade = file_name_parts[1] if len(file_name_parts) > 1 else "Grade"

    # Create PDF object
    pdf = StyledPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # Instructions box
    pdf.set_text_color(91, 44, 250)
    pdf.set_font("Arial", style="B", size=18)
    pdf.ln(2)
    pdf.set_fill_color(240, 240, 255)
    pdf.set_text_color(0, 0, 0)
    left_margin = 10
    right_margin = 10
    box_width = pdf.w - left_margin - right_margin
    
    # First line with bold "Instructions:" - centered
    pdf.set_font("Arial", style="B", size=12)
    bold_text = "Instructions:"
    pdf.set_font("Arial", style="", size=12)
    regular_text = " For each of the following, provide the most accurate and complete response."
    
    # Calculate total width and centering position
    pdf.set_font("Arial", style="B", size=12)
    bold_width = pdf.get_string_width(bold_text)
    pdf.set_font("Arial", style="", size=12)
    regular_width = pdf.get_string_width(regular_text)
    total_width = bold_width + regular_width
    start_x = left_margin + (box_width - total_width) / 2
    
    # Draw the filled background
    pdf.set_x(left_margin)
    pdf.cell(box_width, 8, "", ln=0, fill=True)
    
    # Position and write bold text
    pdf.set_xy(start_x, pdf.get_y())
    pdf.set_font("Arial", style="B", size=12)
    pdf.cell(bold_width, 8, bold_text, ln=0)
    
    # Write regular text
    pdf.set_font("Arial", style="", size=12)
    pdf.cell(regular_width, 8, regular_text, ln=1)
    
    # Second line - centered
    pdf.set_x(left_margin)
    pdf.cell(box_width, 8, "Explanations are at the end.", ln=1, fill=True, align='C')
    pdf.ln(6)

    import re
    section_headers = ["Knowledge", "Thinking", "Application", "Communication"]
    
    # More flexible section detection - matches any Section X with keywords
    def parse_sections(text, section_headers_list):
        section_map = {h: [] for h in section_headers_list}
        current_section = None
        current_question = []
        
        for line in text.splitlines():
            line_strip = line.strip()
            
            # Check if this line is a section header
            if line_strip.startswith("**Section") or line_strip.startswith("**section"):
                # Save previous question if exists
                if current_question and current_section:
                    section_map[current_section].append("\n".join(current_question))
                    current_question = []
                
                # Try to map to our standard headers based on keywords
                line_lower = line_strip.lower()
                if "knowledge" in line_lower or "understanding" in line_lower:
                    current_section = "Knowledge"
                elif "thinking" in line_lower or "inquiry" in line_lower:
                    current_section = "Thinking"
                elif "application" in line_lower or "making connections" in line_lower:
                    current_section = "Application"
                elif "communication" in line_lower:
                    current_section = "Communication"
                continue
            
            # Skip markdown markers, dividers, and empty lines
            if line_strip.startswith("**") or line_strip.startswith("---") or not line_strip:
                continue
            
            # Add content to current section
            if current_section:
                # Check if this is a new question (starts with number)
                if re.match(r'^\d+\.', line_strip):
                    # Save previous question
                    if current_question:
                        section_map[current_section].append("\n".join(current_question))
                    current_question = [line_strip]
                # Check if this is a multiple choice option
                elif re.match(r'^\s*[a-e]\)', line_strip):
                    current_question.append(line_strip)
                # Otherwise, it's a continuation of the current question
                else:
                    if current_question:
                        current_question.append(line_strip)
                    else:
                        current_question = [line_strip]
        
        # Don't forget the last question
        if current_question and current_section:
            section_map[current_section].append("\n".join(current_question))
        
        return section_map

    # Find where answers start
    answers_start = None
    lines = questions.splitlines()
    for idx, line in enumerate(lines):
        line_lower = line.strip().lower()
        if line_lower.startswith("**answer") or line_lower.startswith("**solution") or "end of test" in line_lower:
            answers_start = idx
            break

    if answers_start is not None:
        questions_text = "\n".join(lines[:answers_start])
        answers_text = "\n".join(lines[answers_start:])
    else:
        questions_text = questions
        answers_text = ""

    question_map = parse_sections(questions_text, section_headers)

    # Render questions under each section header
    question_number = 1
    for section in section_headers:
        if question_map[section]:  # Only show section if it has content
            pdf.ln(12)
            pdf.set_font("Arial", style="B", size=15)  # Bold section headers
            pdf.cell(0, 10, section + ":", ln=1, fill=True)
            pdf.set_font("Arial", size=12)
            for q in question_map[section]:
                # Each q is a complete question with possible multiple choice options
                lines_in_q = q.split('\n')
                
                # First line is the question text (remove old numbering)
                first_line = re.sub(r'^\d+\.\s*', '', lines_in_q[0].strip())
                pdf.multi_cell(0, 10, f"{question_number}. {first_line}")
                
                # Remaining lines are multiple choice options or continuation
                for line in lines_in_q[1:]:
                    line_clean = line.strip()
                    if line_clean:
                        # Check if it's a multiple choice option
                        if re.match(r'^[a-e]\)', line_clean):
                            pdf.ln(1)
                            pdf.set_x(pdf.l_margin + 5)  # Indent choices
                            pdf.multi_cell(0, 8, line_clean)
                        else:
                            # It's a continuation of the question
                            pdf.multi_cell(0, 10, line_clean)
                
                pdf.ln(5)  # More vertical space between questions
                question_number += 1

    # Parse and render answers if present
    answer_map = {h: [] for h in section_headers}
    if answers_text:
        print(f"DEBUG: Parsing answers from {len(answers_text)} characters")
        # Use same section parsing for answers
        answer_map = parse_sections(answers_text, section_headers)
        for sec, items in answer_map.items():
            print(f"DEBUG: Answer section '{sec}' has {len(items)} items")

    # Add Solutions page
    pdf.add_page()
    pdf.set_font("Arial", style="B", size=20)
    pdf.set_text_color(91, 44, 250)
    pdf.cell(0, 10, "Solutions", ln=True, align="C")
    pdf.ln(10)
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Arial", size=12)
    
    answer_number = 1
    for section in section_headers:
        if answer_map[section]:
            pdf.set_font("Arial", style="B", size=15)
            pdf.cell(0, 10, section + ":", ln=1, fill=True)
            pdf.set_font("Arial", size=12)
            for a in answer_map[section]:
                # Each answer might have multiple lines
                lines_in_a = a.split('\n')
                
                # First line is the answer text (remove old numbering)
                first_line = re.sub(r'^\d+\.\s*', '', lines_in_a[0].strip())
                pdf.multi_cell(0, 10, f"{answer_number}. {first_line}")
                
                # Remaining lines are continuation of the answer
                for line in lines_in_a[1:]:
                    line_clean = line.strip()
                    if line_clean and not line_clean.startswith("*Note"):
                        pdf.multi_cell(0, 10, line_clean)
                
                pdf.ln(5)  # Same vertical spacing as questions
                answer_number += 1

    original_file_name = os.path.splitext(object_key)[0]
    output_pdf_path = f"{original_file_name}_Practice_Test.pdf"
    pdf.output(output_pdf_path)
    print(f"Practice test saved to '{output_pdf_path}'.")

    second_s3_client = boto3.client(
        's3',
        aws_access_key_id=second_aws_access_key_id,
        aws_secret_access_key=second_aws_secret_access_key
    )

    try:
        second_s3_client.upload_file(
            Filename=output_pdf_path,
            Bucket=second_s3_bucket_name,
            Key=os.path.basename(output_pdf_path)
        )
        print(f"PDF successfully uploaded to S3 bucket '{second_s3_bucket_name}' as '{os.path.basename(output_pdf_path)}'.")
    except Exception as e:
        logging.error(f"Error uploading PDF to S3: {e}")

    try:
        os.remove(output_pdf_path)
        os.remove(object_key)
        os.remove(docx_output_file_path)
        print("Temporary files deleted successfully.")
    except Exception as e:
        logging.error(f"Error deleting temporary files: {e}")