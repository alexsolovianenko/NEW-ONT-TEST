import boto3
import os
from dotenv import load_dotenv
import logging
from datetime import datetime
from docx import Document
from fpdf import FPDF, HTMLMixin

load_dotenv()
second_aws_access_key_id = os.getenv('SECOND_AWS_ACCESS_KEY_ID')
second_aws_secret_access_key = os.getenv('SECOND_AWS_SECRET_ACCESS_KEY')
second_aws_region = os.getenv('SECOND_AWS_REGION')
second_s3_bucket_name = os.getenv('SECOND_S3_BUCKET_NAME')




class StyledPDF(FPDF, HTMLMixin):
    def header(self):
        self.set_fill_color(91, 44, 250)  
        self.set_text_color(255, 255, 255)  
        self.set_font("Arial", style="B", size=10)  # Use default font
        self.set_text_color(0, 0, 0)
        self.set_xy(-40, 5)
        self.cell(30, 10, f"Page {self.page_no()}", align="R")

        # Title
        self.set_xy(10, 18)
        self.set_text_color(255, 255, 255)
        self.set_font("Arial", style="B", size=46)  # Use default font
        self.cell(0, 28, "Ontario Tests", ln=True, align="C", fill=True)
        self.ln(5)

    def footer(self):
        from datetime import datetime
        self.set_y(-15)
        self.set_font("Arial", style="", size=11)  # Use default font
        self.set_text_color(138, 153, 163)  # #8a99a3
        dev_notice = f"© {datetime.now().year} Ontario Tests | Early Development Version"
        self.cell(0, 10, dev_notice, align="C")
        self.set_text_color(150, 150, 150)

if __name__ == "__main__":
    # Load questions from the updated practice_test.txt file
    practice_test_path = r"C:\Users\solov\ontario test\practice_test.txt"
    if not os.path.exists(practice_test_path):
        print(f"Error: The file '{practice_test_path}' does not exist.")
        exit(1)

    # Function to clean text for latin-1 encoding
    def clean_text_for_pdf(text):
        """Replace unicode characters that can't be encoded in latin-1"""
        replacements = {
            '\u2019': "'",  # Right single quotation mark
            '\u2018': "'",  # Left single quotation mark
            '\u201c': '"',  # Left double quotation mark
            '\u201d': '"',  # Right double quotation mark
            '\u2013': '-',  # En dash
            '\u2014': '-',  # Em dash
            '\u2026': '...',  # Ellipsis
        }
        for old, new in replacements.items():
            text = text.replace(old, new)
        return text

    # Load questions from file
    with open(practice_test_path, encoding="utf-8") as file:
        questions = clean_text_for_pdf(file.read())

    # Create PDF object
    pdf = StyledPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # Instructions box
    pdf.set_text_color(91, 44, 250)
    pdf.set_font("Arial", style="B", size=18)
    pdf.ln(2)
    pdf.set_fill_color(240, 240, 255)
    pdf.set_text_color(0, 0, 0)
    left_margin = 10
    right_margin = 10
    box_width = pdf.w - left_margin - right_margin
    
    # First line with bold "Instructions:" - centered
    pdf.set_font("Arial", style="B", size=12)
    bold_text = "Instructions:"
    pdf.set_font("Arial", style="", size=12)
    regular_text = " For each of the following, provide the most accurate and complete response."
    
    # Calculate total width and centering position
    pdf.set_font("Arial", style="B", size=12)
    bold_width = pdf.get_string_width(bold_text)
    pdf.set_font("Arial", style="", size=12)
    regular_width = pdf.get_string_width(regular_text)
    total_width = bold_width + regular_width
    start_x = left_margin + (box_width - total_width) / 2
    
    # Draw the filled background
    pdf.set_x(left_margin)
    pdf.cell(box_width, 8, "", ln=0, fill=True)
    
    # Position and write bold text
    pdf.set_xy(start_x, pdf.get_y())
    pdf.set_font("Arial", style="B", size=12)
    pdf.cell(bold_width, 8, bold_text, ln=0)
    
    # Write regular text
    pdf.set_font("Arial", style="", size=12)
    pdf.cell(regular_width, 8, regular_text, ln=1)
    
    # Second line - centered
    pdf.set_x(left_margin)
    pdf.cell(box_width, 8, "Explanations are at the end.", ln=1, fill=True, align='C')
    pdf.ln(6)


    import re
    section_headers = ["Knowledge", "Thinking", "Application", "Communication"]
    
    def parse_sections(text):
        section_map = {h: [] for h in section_headers}
        current_section = None
        current_question = []
        
        for line in text.splitlines():
            line_strip = line.strip()
            
            # Skip empty lines and dividers
            if not line_strip or line_strip == '---':
                continue
            
            # Skip title lines that start with #
            if line_strip.startswith('# Practice Test') or line_strip.startswith('# **Solutions**'):
                continue
            
            # Check for section headers (## **Knowledge**, etc.)
            if '**Knowledge' in line_strip and line_strip.startswith('##'):
                if current_question and current_section:
                    section_map[current_section].append(current_question)
                    current_question = []
                current_section = "Knowledge"
                continue
            elif '**Thinking' in line_strip and line_strip.startswith('##'):
                if current_question and current_section:
                    section_map[current_section].append(current_question)
                    current_question = []
                current_section = "Thinking"
                continue
            elif '**Application' in line_strip and line_strip.startswith('##'):
                if current_question and current_section:
                    section_map[current_section].append(current_question)
                    current_question = []
                current_section = "Application"
                continue
            elif '**Communication' in line_strip and line_strip.startswith('##'):
                if current_question and current_section:
                    section_map[current_section].append(current_question)
                    current_question = []
                current_section = "Communication"
                continue
            
            # Add content to current section
            if current_section:
                # Check if this is a new question (starts with **number.**)
                if re.match(r'\*\*\d+\.\*\*', line_strip):
                    # Save previous question
                    if current_question:
                        section_map[current_section].append(current_question)
                    # Start new question - remove the **number.** part
                    question_text = re.sub(r'\*\*\d+\.\*\*\s*', '', line_strip)
                    current_question = [question_text]
                # Check if this is a multiple choice option
                elif re.match(r'^[a-d]\)', line_strip):
                    if current_question:
                        current_question.append(line_strip)
                # Otherwise it's continuation of question text
                elif current_question:
                    current_question[0] += " " + line_strip
        
        # Don't forget the last question
        if current_question and current_section:
            section_map[current_section].append(current_question)
        
        return section_map

    # Find where answers start
    answers_start = None
    lines = questions.splitlines()
    for idx, line in enumerate(lines):
        # Look for "# **Solutions**" or similar
        if '**Solutions**' in line.strip() or line.strip().startswith('# **Solutions'):
            answers_start = idx
            break

    if answers_start is not None:
        questions_text = "\n".join(lines[:answers_start])
        answers_text = "\n".join(lines[answers_start:])
    else:
        questions_text = questions
        answers_text = ""

    question_map = parse_sections(questions_text)

    # Render questions under each section header
    question_number = 1
    for section in section_headers:
        if question_map[section]:  # Only show section if it has content
            pdf.ln(12)
            pdf.set_font("Arial", style="B", size=15)
            pdf.set_fill_color(235, 233, 254)
            pdf.cell(0, 10, section + ":", ln=1, fill=True)
            pdf.set_font("Arial", size=12)
            
            for question_parts in question_map[section]:
                # question_parts is a list: [question_text, option_a, option_b, ...]
                if isinstance(question_parts, list):
                    # First item is the question text
                    question_text = question_parts[0]
                    # Remove old numbering
                    question_text = re.sub(r'^\d+\.\s*', '', question_text.strip())
                    pdf.multi_cell(0, 10, f"{question_number}. {question_text}")
                    
                    # Remaining items are multiple choice options
                    if len(question_parts) > 1:
                        pdf.ln(1)
                        for option in question_parts[1:]:
                            pdf.set_x(pdf.l_margin + 5)  # Indent choices
                            pdf.multi_cell(0, 8, option.strip())
                    
                    pdf.ln(5)
                else:
                    # Fallback for plain string
                    question_text = re.sub(r'^\d+\.\s*', '', question_parts.strip())
                    pdf.multi_cell(0, 10, f"{question_number}. {question_text}")
                    pdf.ln(5)
                
                question_number += 1

    # Parse and render answers if present
    answer_map = {h: [] for h in section_headers}
    if answers_text:
        # Parse answers by section
        current_section = None
        
        for line in answers_text.splitlines():
            line_strip = line.strip()
            
            # Skip empty lines and dividers
            if not line_strip or line_strip == '---':
                continue
                
            # Skip "Solutions" header lines
            if '**Solutions**' in line_strip or line_strip.startswith('# **Solutions'):
                continue
            
            # Check for section headers in answers (### Knowledge, ### Thinking, etc.)
            if line_strip.startswith('###'):
                if 'Knowledge' in line_strip:
                    current_section = "Knowledge"
                    continue
                elif 'Thinking' in line_strip:
                    current_section = "Thinking"
                    continue
                elif 'Application' in line_strip:
                    current_section = "Application"
                    continue
                elif 'Communication' in line_strip:
                    current_section = "Communication"
                    continue
            
            # Add answer lines to current section (remove number prefix like "1. ")
            if current_section and line_strip:
                # Remove the number prefix from answers
                answer_text = re.sub(r'^\d+\.\s*', '', line_strip)
                if answer_text:  # Only add if there's text left
                    answer_map[current_section].append(answer_text)

    # Add Solutions page
    pdf.add_page()
    pdf.set_font("Arial", style="B", size=20)
    pdf.set_text_color(91, 44, 250)
    pdf.cell(0, 10, "Solutions", ln=True, align="C")
    pdf.ln(10)
    pdf.set_text_color(0, 0, 0)
    pdf.set_font("Arial", size=12)
    
    for section in section_headers:
        if answer_map[section]:
            pdf.set_font("Arial", style="B", size=15)
            pdf.cell(0, 10, section + ":", ln=1, fill=True)
            pdf.set_font("Arial", size=12)
            section_answer_number = 1
            for a in answer_map[section]:
                # Remove old numbering from answers
                a_clean = re.sub(r'^\d+\.\s*', '', a.strip())
                pdf.multi_cell(0, 10, f"{section_answer_number}. {a_clean}")
                pdf.ln(5)
                section_answer_number += 1

    original_file_name = os.path.splitext(practice_test_path)[0]
    output_pdf_path = "10-science-biology.pdf"
    pdf.output(output_pdf_path)
    print(f"Practice test saved to '{output_pdf_path}'.")

    second_s3_client = boto3.client(
        's3',
        aws_access_key_id=second_aws_access_key_id,
        aws_secret_access_key=second_aws_secret_access_key
    )

    try:
        second_s3_client.upload_file(
            Filename=output_pdf_path,
            Bucket=second_s3_bucket_name,
            Key=os.path.basename(output_pdf_path)
        )
        print(f"PDF successfully uploaded to S3 bucket '{second_s3_bucket_name}' as '{os.path.basename(output_pdf_path)}'.")
    except Exception as e:
        logging.error(f"Error uploading PDF to S3: {e}")

    try:
        os.remove(output_pdf_path)
        print("Temporary files deleted successfully.")
    except Exception as e:
        logging.error(f"Error deleting temporary files: {e}")